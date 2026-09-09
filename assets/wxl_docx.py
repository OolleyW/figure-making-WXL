"""Word assembly for WXL figures.

Turns WXL figures into a manuscript-style ``.docx`` that follows the same paper
contract the figures follow:

  * A4, 10 mm side margins, so the usable text width is exactly 190 mm
  * figures inserted at 100 % of their measured physical width, which keeps the
    10 pt figure text equal to 10 pt Word body text
  * figure captions BELOW the figure
  * table captions ABOVE the table, three-line tables (top / header / bottom)
  * section headings in black CJK serif bold with double line spacing
  * body text in Times New Roman 10 pt, double spacing, 2-character indent

Requires ``python-docx`` (``pip install python-docx``). Import it after adding
the skill's ``assets`` directory to ``sys.path``.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from wxl_style import measure_width_mm

#: Latin face for body text and captions.
TNR = "Times New Roman"
#: CJK face for headings and Chinese body text. Override if the machine lacks it.
CJK_FONT = "SimSun"

__all__ = [
    "TNR", "CJK_FONT", "new_document", "add_title", "add_heading",
    "add_paragraph", "add_caption", "add_figure", "add_figure_block",
    "add_three_line_table", "add_page_break", "save_document",
]


# --------------------------------------------------------------------------
def _set_run(run, font=TNR, size=10.0, bold=False, cjk=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk or CJK_FONT)


def new_document(side_mm: float = 10.0, top_mm: float = 20.0,
                 bottom_mm: float = 20.0, base_size: float = 10.0) -> Document:
    """A4 document whose usable width is ``210 - 2 * side_mm`` millimetres."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = TNR
    style.font.size = Pt(base_size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(side_mm / 10.0)
    sec.top_margin = Cm(top_mm / 10.0)
    sec.bottom_margin = Cm(bottom_mm / 10.0)
    return doc


def add_paragraph(doc, text, size=10.0, bold=False, align=None, font=TNR,
                  spacing=2.0, indent_chars=2, space_after=0.0):
    """Body paragraph: Times New Roman 10 pt, double spacing, 2-char indent."""
    p = doc.add_paragraph()
    _set_run(p.add_run(text), font=font, size=size, bold=bold)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(space_after)
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    return p


def add_heading(doc, text, size=12.0):
    """Section heading: black CJK serif bold, double spacing, no indent."""
    p = doc.add_paragraph()
    _set_run(p.add_run(text), font=CJK_FONT, size=size, bold=True)
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(6)
    pf.space_after = Pt(2)
    return p


def add_title(doc, text, size=16.0, font=None, space_after=10.0):
    """Document title: centered, bold, double spacing, no indent."""
    p = doc.add_paragraph()
    _set_run(p.add_run(text), font=font or CJK_FONT, size=size, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(space_after)
    return p


def add_caption(doc, text, size=10.0, spacing=1.5, space_after=6.0):
    """Centered caption. Use BELOW a figure and ABOVE a table."""
    return add_paragraph(doc, text, size=size, align=WD_ALIGN_PARAGRAPH.CENTER,
                         font=TNR, spacing=spacing, indent_chars=0,
                         space_after=space_after)


def add_page_break(doc):
    doc.add_page_break()


def save_document(doc, path):
    """Save ``doc`` to ``path``, falling back to a timestamped name if locked.

    Word keeps an exclusive lock on an open document, so re-running a build
    script while the previous output is open raises ``PermissionError``. In that
    case the document is written next to it with a timestamp suffix and the path
    actually used is returned.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(path)
        return path
    except PermissionError:
        stamped = path.with_name(
            f"{path.stem}_{datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        doc.save(stamped)
        print(f"note: {path.name} is locked (open in Word?), wrote {stamped.name}")
        return stamped


def add_figure(doc, png, dpi: int = 600, space_after: float = 4.0) -> float:
    """Insert ``png`` at 100 % of its physical size, centered.

    Returns the inserted width in millimetres. The image must have been produced
    by :func:`wxl_style.finalize_figure` with ``target_width_mm`` so that its
    physical width equals the intended column width.
    """
    mm = measure_width_mm(png, dpi)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(space_after)
    p.add_run().add_picture(str(png), width=Cm(mm / 10.0))
    return mm


def add_figure_block(doc, png, caption, dpi: int = 600) -> float:
    """Figure + caption below it. Returns the inserted width in millimetres."""
    mm = add_figure(doc, png, dpi=dpi)
    add_caption(doc, caption)
    return mm


# --------------------------------------------------------------------------
def _cell_borders(cell, top=None, bottom=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, sz in (("top", top), ("bottom", bottom)):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))          # eighths of a point
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
    for edge in ("left", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def add_three_line_table(doc, headers, rows, col_widths_cm=None,
                         caption=None, font_size=10.0, spacing=1.5):
    """Three-line table: 1.5 pt top rule, 0.75 pt header rule, 1.5 pt bottom rule.

    ``caption`` is placed ABOVE the table, centered. Column widths default to an
    even split of the usable width (190 mm with the default margins).
    """
    if caption:
        add_caption(doc, caption, size=font_size)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def fill(cells, values, bold):
        for cell, value in zip(cells, values):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = spacing
            p.paragraph_format.space_after = Pt(0)
            _set_run(p.add_run(str(value)), size=font_size, bold=bold)

    fill(table.rows[0].cells, headers, True)
    for i, row in enumerate(rows, start=1):
        fill(table.rows[i].cells, row, False)

    for i, row in enumerate(table.rows):
        for cell in row.cells:
            top = 12 if i == 0 else None
            bottom = 6 if i == 0 else (12 if i == len(table.rows) - 1 else None)
            _cell_borders(cell, top=top, bottom=bottom)

    if col_widths_cm is None:
        col_widths_cm = [19.0 / len(headers)] * len(headers)
    for row in table.rows:
        for cell, w in zip(row.cells, col_widths_cm):
            cell.width = Cm(w)
    return table
